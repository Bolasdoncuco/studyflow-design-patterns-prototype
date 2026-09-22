"""Build the English submission documents from the Markdown source files."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "submission"

def add_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(set(cell) <= {"-", ":", " "} for cell in cells):
                    rows.append(cells)
                index += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for col_index, value in enumerate(row):
                        table.cell(row_index, col_index).text = value
                doc.add_paragraph()
            continue
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:])
        else:
            p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(6)
        index += 1

def build(source: str, target: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    add_markdown(doc, (OUT / source).read_text(encoding="utf-8"))
    doc.save(OUT / target)

if __name__ == "__main__":
    build("PART_A_ANALYSIS.md", "Part_A_Analysis.docx")
    build("STUDYFLOW_JUSTIFICATION.md", "StudyFlow_Justification.docx")
